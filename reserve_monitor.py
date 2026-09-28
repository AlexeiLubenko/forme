#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Daily reserve monitoring pipeline.

Main outputs:
- processed daily contract snapshot (CSV + Parquet)
- day-to-day contract movements (CSV + Parquet)
- cumulative Qlik-ready marts
- daily Excel monitoring report
- data-quality JSON/CSV

The script does NOT assume that changes in actual reserve are mathematically
explained by EAD*PD*LGD. Exact reconciliation is performed at the top level:
NEW + EXIT + EXISTING = actual day-to-day reserve change.
Inside EXISTING, driver classification is rule-based and auditable.
"""

from __future__ import annotations

import argparse
import json
import logging
import re
import sys
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

import numpy as np
import pandas as pd
import yaml
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# ------------------------------ constants ------------------------------
TRUE_VALUES = {"1", "true", "yes", "y", "да", "истина", "x", "+"}
FALSE_VALUES = {"0", "false", "no", "n", "нет", "ложь", "", "nan", "none"}
MISSING_KEY = "__MISSING__"
NO_TRANCHE = "__NO_TRANCHE__"


@dataclass
class RunContext:
    cfg: dict[str, Any]
    cfg_path: Path
    project_root: Path
    raw_dir: Path
    processed_daily_dir: Path
    processed_movements_dir: Path
    qlik_dir: Path
    reports_dir: Path
    logs_dir: Path


# ------------------------------ helpers ------------------------------
def read_yaml(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as f:
        return yaml.safe_load(f)


def resolve_context(cfg_path: Path) -> RunContext:
    cfg_path = cfg_path.resolve()
    cfg = read_yaml(cfg_path)
    project_root = cfg_path.parent.resolve()

    def p(name: str) -> Path:
        value = Path(cfg["paths"][name])
        return value if value.is_absolute() else project_root / value

    ctx = RunContext(
        cfg=cfg,
        cfg_path=cfg_path,
        project_root=project_root,
        raw_dir=p("raw_dir"),
        processed_daily_dir=p("processed_daily_dir"),
        processed_movements_dir=p("processed_movements_dir"),
        qlik_dir=p("qlik_dir"),
        reports_dir=p("reports_dir"),
        logs_dir=p("logs_dir"),
    )
    for folder in (
        ctx.raw_dir,
        ctx.processed_daily_dir,
        ctx.processed_movements_dir,
        ctx.qlik_dir,
        ctx.reports_dir,
        ctx.logs_dir,
    ):
        folder.mkdir(parents=True, exist_ok=True)
    return ctx


def setup_logging(ctx: RunContext) -> None:
    log_path = ctx.logs_dir / "reserve_monitor.log"
    handlers = [logging.StreamHandler(sys.stdout), logging.FileHandler(log_path, encoding="utf-8")]
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s | %(levelname)s | %(message)s",
        handlers=handlers,
        force=True,
    )


def normalize_header(value: Any) -> str:
    """Remove numeric prefixes like '45. EAD', normalize spaces and NBSP."""
    s = "" if value is None else str(value)
    s = s.replace("\ufeff", "").replace("\xa0", " ").replace("ё", "е")
    s = re.sub(r"^\s*\d+\s*[\.)]\s*", "", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


def normalize_headers(df: pd.DataFrame) -> pd.DataFrame:
    result = df.copy()
    result.columns = [normalize_header(c) for c in result.columns]
    return result


def safe_str(series: pd.Series, missing: str = MISSING_KEY) -> pd.Series:
    out = series.astype("string").fillna(missing).str.strip()
    out = out.replace({"": missing, "<NA>": missing, "nan": missing, "None": missing})
    return out


def as_numeric(series: pd.Series) -> pd.Series:
    if pd.api.types.is_numeric_dtype(series):
        return pd.to_numeric(series, errors="coerce")
    cleaned = (
        series.astype("string")
        .str.replace("\xa0", "", regex=False)
        .str.replace(" ", "", regex=False)
        .str.replace(",", ".", regex=False)
    )
    return pd.to_numeric(cleaned, errors="coerce")


def as_date(series: pd.Series) -> pd.Series:
    # format='mixed' correctly handles ISO and local DD.MM.YYYY values in pandas >=2.
    return pd.to_datetime(series, errors="coerce", format="mixed", dayfirst=True)


def normalize_stage(series: pd.Series) -> pd.Series:
    s = series.astype("string").fillna("UNKNOWN").str.upper().str.strip()
    mapped = pd.Series("UNKNOWN", index=s.index, dtype="string")
    mapped[s.str.contains(r"(?:^|\D)1(?:$|\D)|STAGE\s*1|СТАДИ[ЯИ]\s*1", regex=True, na=False)] = "1"
    mapped[s.str.contains(r"(?:^|\D)2(?:$|\D)|STAGE\s*2|СТАДИ[ЯИ]\s*2", regex=True, na=False)] = "2"
    mapped[s.str.contains(r"(?:^|\D)3(?:$|\D)|STAGE\s*3|СТАДИ[ЯИ]\s*3", regex=True, na=False)] = "3"
    numeric = pd.to_numeric(s, errors="coerce")
    for value in (1, 2, 3):
        mapped[numeric.eq(value)] = str(value)
    return mapped


def normalize_bool(series: pd.Series) -> pd.Series:
    s = series.astype("string").fillna("").str.lower().str.strip()
    result = pd.Series(False, index=s.index)
    result[s.isin(TRUE_VALUES)] = True
    return result


def relative_change(new: pd.Series, old: pd.Series) -> pd.Series:
    denominator = old.abs().replace(0, np.nan)
    return (new - old) / denominator


def first_existing(columns: Iterable[str], df: pd.DataFrame) -> list[str]:
    return [c for c in columns if c in df.columns]


def read_input(path: Path, sheet_name: Any = 0) -> pd.DataFrame:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm", ".xls"}:
        return pd.read_excel(path, sheet_name=sheet_name)
    if suffix == ".csv":
        try:
            return pd.read_csv(path, sep=None, engine="python")
        except UnicodeDecodeError:
            return pd.read_csv(path, sep=None, engine="python", encoding="cp1251")
    if suffix == ".parquet":
        return pd.read_parquet(path)
    raise ValueError(f"Unsupported input format: {path}")


def find_latest_raw(ctx: RunContext) -> Path:
    candidates = [
        p for p in ctx.raw_dir.iterdir()
        if p.is_file() and p.suffix.lower() in {".xlsx", ".xlsm", ".xls", ".csv", ".parquet"}
    ]
    if not candidates:
        raise FileNotFoundError(f"No raw files found in {ctx.raw_dir}")
    return max(candidates, key=lambda p: p.stat().st_mtime)


def validate_required_columns(df: pd.DataFrame, cfg: dict[str, Any]) -> None:
    c = cfg["columns"]
    required = {
        c["report_date"], c["crm_id"], c["contract_number"], c["stage"], c["ead"],
        c["source"], c["fx_rate"], c["currency"], c["pd_pit"], c["lgd"],
    }
    required.update(cfg["reserve_columns"])
    missing = sorted(x for x in required if x not in df.columns)
    if missing:
        raise ValueError(
            "Input file is missing required columns after header normalization:\n- " + "\n- ".join(missing)
        )


def get_single_report_date(df: pd.DataFrame, report_date_col: str) -> pd.Timestamp:
    dates = as_date(df[report_date_col]).dropna().dt.normalize().unique()
    if len(dates) != 1:
        readable = [str(pd.Timestamp(d).date()) for d in dates[:10]]
        raise ValueError(f"Expected exactly one report date per snapshot, got {len(dates)}: {readable}")
    return pd.Timestamp(dates[0]).normalize()


def make_contract_key(df: pd.DataFrame, cfg: dict[str, Any]) -> pd.Series:
    key_cols = cfg["input"]["key_columns"]
    missing = [c for c in key_cols if c not in df.columns]
    if missing:
        raise ValueError(f"Key columns missing: {missing}")
    pieces: list[pd.Series] = []
    for col in key_cols:
        missing_marker = NO_TRANCHE if col == cfg["columns"]["tranche_number"] else MISSING_KEY
        pieces.append(safe_str(df[col], missing=missing_marker))
    result = pieces[0]
    for piece in pieces[1:]:
        result = result.str.cat(piece, sep="|")
    return result


def ensure_no_duplicate_keys(df: pd.DataFrame, cfg: dict[str, Any]) -> None:
    duplicated = df["contract_key"].duplicated(keep=False)
    n = int(duplicated.sum())
    if n == 0:
        return
    policy = str(cfg["input"].get("duplicate_policy", "error")).lower()
    sample = df.loc[duplicated, ["contract_key"]].head(10).to_dict("records")
    if policy == "error":
        raise ValueError(f"Found {n} rows with duplicated contract_key. Sample: {sample}")
    logging.warning("Found %s duplicated rows. duplicate_policy=%s; keeping first row per key.", n, policy)
    df.drop_duplicates("contract_key", keep="first", inplace=True)


def prepare_snapshot(raw: pd.DataFrame, cfg: dict[str, Any]) -> tuple[pd.DataFrame, pd.Timestamp]:
    df = normalize_headers(raw)
    validate_required_columns(df, cfg)
    c = cfg["columns"]
    report_date = get_single_report_date(df, c["report_date"])

    df[c["report_date"]] = as_date(df[c["report_date"]]).dt.normalize()
    df["contract_key"] = make_contract_key(df, cfg)
    ensure_no_duplicate_keys(df, cfg)

    # Numeric model/risk fields.
    numeric_candidates = [
        c["ead"], c["fx_rate"], c["pd_contract"], c["pd_pit"], c["pd_lt_max"],
        c["macro_addon"], c["ccf"], c["lgd"], c["lgd_lr"], c["lgd_id"], c["lgd_alt"],
        c["overdue_days_contract"], c["overdue_days_client"],
    ] + cfg["reserve_columns"]
    for col in first_existing(numeric_candidates, df):
        df[col] = as_numeric(df[col])

    # Technical normalized dimensions.
    df["stage_norm"] = normalize_stage(df[c["stage"]])
    df["individual_flag"] = normalize_bool(df[c["individual_assessment"]]) if c["individual_assessment"] in df else False
    df["rating_norm"] = safe_str(df[c["rating"]], missing="UNKNOWN") if c["rating"] in df else "UNKNOWN"
    df["overdue_bucket_norm"] = safe_str(df[c["overdue_bucket"]], missing="UNKNOWN") if c["overdue_bucket"] in df else "UNKNOWN"
    df["currency_norm"] = safe_str(df[c["currency"]], missing="UNKNOWN")

    # Reserve components and EAD.
    reserve_cols = first_existing(cfg["reserve_columns"], df)
    overdue_cols = first_existing(cfg.get("overdue_reserve_columns", []), df)
    df["reserve_original"] = df[reserve_cols].fillna(0).sum(axis=1)
    df["reserve_overdue_original"] = df[overdue_cols].fillna(0).sum(axis=1) if overdue_cols else 0.0
    df["ead_original"] = df[c["ead"]].fillna(0)

    values_in_byn = bool(cfg["input"].get("monetary_values_in_byn", True))
    fx = df[c["fx_rate"]].fillna(1.0).replace(0, np.nan).fillna(1.0)
    if values_in_byn:
        df["ead_byn"] = df["ead_original"]
        df["reserve_byn"] = df["reserve_original"]
        df["reserve_overdue_byn"] = df["reserve_overdue_original"]
    else:
        df["ead_byn"] = df["ead_original"] * fx
        df["reserve_byn"] = df["reserve_original"] * fx
        df["reserve_overdue_byn"] = df["reserve_overdue_original"] * fx

    df["ead_mln"] = df["ead_byn"] / 1_000_000.0
    df["reserve_mln"] = df["reserve_byn"] / 1_000_000.0
    df["reserve_overdue_mln"] = df["reserve_overdue_byn"] / 1_000_000.0
    df["reserve_rate"] = np.where(df["ead_byn"].ne(0), df["reserve_byn"] / df["ead_byn"], np.nan)
    df["report_date"] = report_date

    # Friendly aliases used by marts/Qlik.
    alias_map = {
        c["crm_id"]: "crm_id",
        c["unp"]: "unp",
        c["client_name"]: "client_name",
        c["source"]: "source",
        c["contract_number"]: "contract_number",
        c["tranche_number"]: "tranche_number",
        c["industry"]: "industry",
        c["subportfolio"]: "subportfolio",
        c["business_segment"]: "business_segment",
        c["risk_segment"]: "risk_segment",
        c["rating"]: "rating",
        c["overdue_bucket"]: "overdue_bucket",
        c["overdue_days_contract"]: "overdue_days_contract",
        c["default_date"]: "default_date",
        c["pd_pit"]: "pd_pit",
        c["pd_contract"]: "pd_contract",
        c["lgd"]: "lgd",
        c["macro_addon"]: "macro_addon",
        c["ccf"]: "ccf",
        c["fx_rate"]: "fx_rate",
    }
    for source_col, target_col in alias_map.items():
        if source_col in df.columns:
            df[target_col] = df[source_col]
    if "default_date" in df.columns:
        df["default_date"] = as_date(df["default_date"])
    if "overdue_days_contract" in df.columns:
        df["overdue_days_contract"] = as_numeric(df["overdue_days_contract"])

    return df, report_date


def select_contract_daily_columns(df: pd.DataFrame) -> pd.DataFrame:
    preferred = [
        "report_date", "contract_key", "crm_id", "unp", "client_name", "source",
        "contract_number", "tranche_number", "industry", "subportfolio", "business_segment",
        "risk_segment", "stage_norm", "rating_norm", "overdue_bucket_norm", "currency_norm",
        "individual_flag", "default_date", "overdue_days_contract", "ead_original", "ead_byn", "ead_mln", "reserve_original",
        "reserve_byn", "reserve_mln", "reserve_overdue_mln", "reserve_rate", "pd_contract",
        "pd_pit", "lgd", "macro_addon", "ccf", "fx_rate",
    ]
    return df[[c for c in preferred if c in df.columns]].copy()


def save_frame(df: pd.DataFrame, csv_path: Path, parquet_path: Path | None = None) -> None:
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(csv_path, index=False, encoding="utf-8-sig")
    if parquet_path is not None:
        parquet_path.parent.mkdir(parents=True, exist_ok=True)
        try:
            df.to_parquet(parquet_path, index=False)
        except (ImportError, ModuleNotFoundError):
            logging.warning("Parquet engine is unavailable; CSV output will be used instead: %s", csv_path)


def previous_daily_file(ctx: RunContext, current_date: pd.Timestamp) -> Path | None:
    candidates: list[tuple[pd.Timestamp, Path]] = []
    for pattern in ("*.parquet", "*.csv"):
        for p in ctx.processed_daily_dir.glob(pattern):
            try:
                d = pd.Timestamp(p.stem).normalize()
            except Exception:
                continue
            if d < current_date:
                candidates.append((d, p))
    if not candidates:
        return None
    # Prefer parquet when both formats exist for the same date.
    candidates.sort(key=lambda x: (x[0], x[1].suffix == ".parquet"))
    return candidates[-1][1]


def classify_rating(prev: pd.Series, curr: pd.Series, rating_order: list[str]) -> pd.Series:
    prev = prev.astype("string").fillna("UNKNOWN")
    curr = curr.astype("string").fillna("UNKNOWN")
    result = pd.Series("UNCHANGED", index=prev.index, dtype="string")
    changed = prev.ne(curr)
    if not rating_order:
        result[changed] = "RATING_CHANGED"
        return result
    rank = {str(v): i for i, v in enumerate(rating_order)}
    prev_rank = prev.map(rank)
    curr_rank = curr.map(rank)
    result[changed & prev_rank.notna() & curr_rank.notna() & curr_rank.gt(prev_rank)] = "DOWNGRADE"
    result[changed & prev_rank.notna() & curr_rank.notna() & curr_rank.lt(prev_rank)] = "UPGRADE"
    result[changed & (prev_rank.isna() | curr_rank.isna())] = "RATING_CHANGED"
    return result


def build_movements(today: pd.DataFrame, prev: pd.DataFrame, cfg: dict[str, Any]) -> pd.DataFrame:
    t = select_contract_daily_columns(today)
    p = select_contract_daily_columns(prev)
    cmp = t.merge(p, on="contract_key", how="outer", suffixes=("_curr", "_prev"), indicator=True)

    cmp["movement_type"] = cmp["_merge"].map({"left_only": "NEW", "right_only": "EXIT", "both": "EXISTING"}).astype("string")
    cmp["report_date"] = cmp.get("report_date_curr").combine_first(cmp.get("report_date_prev"))
    cmp["previous_report_date"] = cmp.get("report_date_prev")

    for base in ["ead_mln", "reserve_mln", "reserve_overdue_mln", "pd_pit", "lgd", "fx_rate"]:
        curr = pd.to_numeric(cmp.get(f"{base}_curr"), errors="coerce") if f"{base}_curr" in cmp else pd.Series(np.nan, index=cmp.index)
        prevv = pd.to_numeric(cmp.get(f"{base}_prev"), errors="coerce") if f"{base}_prev" in cmp else pd.Series(np.nan, index=cmp.index)
        cmp[f"delta_{base}"] = curr.fillna(0) - prevv.fillna(0)

    # Exact reserve effect by contract.
    cmp["delta_reserve_mln"] = cmp["delta_reserve_mln"]
    cmp["delta_ead_mln"] = cmp["delta_ead_mln"]
    cmp["reserve_effect_mln"] = cmp["delta_reserve_mln"]

    # Migrations/changes.
    stage_prev = normalize_stage(cmp.get("stage_norm_prev", pd.Series("UNKNOWN", index=cmp.index)))
    stage_curr = normalize_stage(cmp.get("stage_norm_curr", pd.Series("UNKNOWN", index=cmp.index)))
    cmp["stage_prev"] = stage_prev
    cmp["stage_curr"] = stage_curr
    cmp["stage_migration"] = np.where(cmp["movement_type"].eq("EXISTING"), stage_prev + "→" + stage_curr, cmp["movement_type"])

    rating_prev = cmp.get("rating_norm_prev", pd.Series("UNKNOWN", index=cmp.index)).fillna("UNKNOWN")
    rating_curr = cmp.get("rating_norm_curr", pd.Series("UNKNOWN", index=cmp.index)).fillna("UNKNOWN")
    cmp["rating_prev"] = rating_prev
    cmp["rating_curr"] = rating_curr
    cmp["rating_migration"] = classify_rating(rating_prev, rating_curr, cfg.get("rating_order", []))

    overdue_prev = cmp.get("overdue_bucket_norm_prev", pd.Series("UNKNOWN", index=cmp.index)).fillna("UNKNOWN")
    overdue_curr = cmp.get("overdue_bucket_norm_curr", pd.Series("UNKNOWN", index=cmp.index)).fillna("UNKNOWN")
    cmp["overdue_prev"] = overdue_prev
    cmp["overdue_curr"] = overdue_curr
    existing = cmp["movement_type"].eq("EXISTING")
    cmp["overdue_changed"] = overdue_prev.ne(overdue_curr) & existing

    # Direction of overdue change: primarily numeric days; bucket order is fallback.
    od_prev = pd.to_numeric(cmp.get("overdue_days_contract_prev"), errors="coerce")
    od_curr = pd.to_numeric(cmp.get("overdue_days_contract_curr"), errors="coerce")
    bucket_order = {str(v): i for i, v in enumerate(cfg.get("overdue_bucket_order", []))}
    prev_bucket_rank = overdue_prev.astype(str).map(bucket_order)
    curr_bucket_rank = overdue_curr.astype(str).map(bucket_order)
    cmp["flag_overdue_deterioration"] = existing & (
        (od_curr.gt(od_prev)) | (od_prev.isna() & curr_bucket_rank.gt(prev_bucket_rank)) |
        (od_prev.notna() & od_curr.notna() & od_curr.gt(0) & od_prev.eq(0))
    )
    cmp["flag_overdue_improvement"] = existing & (
        (od_curr.lt(od_prev)) | (od_prev.isna() & curr_bucket_rank.lt(prev_bucket_rank))
    )

    default_prev = pd.to_datetime(cmp.get("default_date_prev"), errors="coerce")
    default_curr = pd.to_datetime(cmp.get("default_date_curr"), errors="coerce")
    cmp["flag_new_default"] = existing & default_prev.isna() & default_curr.notna()

    thresholds = cfg["thresholds"]
    pd_prev = pd.to_numeric(cmp.get("pd_pit_prev"), errors="coerce")
    pd_curr = pd.to_numeric(cmp.get("pd_pit_curr"), errors="coerce")
    lgd_prev = pd.to_numeric(cmp.get("lgd_prev"), errors="coerce")
    lgd_curr = pd.to_numeric(cmp.get("lgd_curr"), errors="coerce")
    ead_prev = pd.to_numeric(cmp.get("ead_mln_prev"), errors="coerce")
    ead_curr = pd.to_numeric(cmp.get("ead_mln_curr"), errors="coerce")
    fx_prev = pd.to_numeric(cmp.get("fx_rate_prev"), errors="coerce")
    fx_curr = pd.to_numeric(cmp.get("fx_rate_curr"), errors="coerce")

    cmp["pd_rel_change"] = relative_change(pd_curr, pd_prev)
    cmp["lgd_abs_change"] = lgd_curr - lgd_prev
    cmp["ead_rel_change"] = relative_change(ead_curr, ead_prev)
    cmp["fx_rel_change"] = relative_change(fx_curr, fx_prev)

    cmp["flag_stage_change"] = existing & stage_prev.ne(stage_curr)
    cmp["flag_stage_deterioration"] = existing & (
        (stage_prev.eq("1") & stage_curr.isin(["2", "3"])) | (stage_prev.eq("2") & stage_curr.eq("3"))
    )
    cmp["flag_stage_improvement"] = existing & (
        (stage_prev.eq("3") & stage_curr.isin(["1", "2"])) | (stage_prev.eq("2") & stage_curr.eq("1"))
    )
    cmp["flag_rating_change"] = existing & cmp["rating_migration"].ne("UNCHANGED")
    cmp["flag_rating_downgrade"] = existing & cmp["rating_migration"].eq("DOWNGRADE")
    cmp["flag_overdue_change"] = existing & cmp["overdue_changed"]
    cmp["flag_pd_change"] = existing & (
        cmp["pd_rel_change"].abs().ge(float(thresholds["pd_relative_change"])) |
        (pd_curr - pd_prev).abs().ge(float(thresholds["pd_absolute_change"]))
    )
    cmp["flag_lgd_change"] = existing & cmp["lgd_abs_change"].abs().ge(float(thresholds["lgd_absolute_change"]))
    cmp["flag_ead_change"] = existing & (
        cmp["ead_rel_change"].abs().ge(float(thresholds["ead_relative_change"])) |
        cmp["delta_ead_mln"].abs().ge(float(thresholds["ead_absolute_change_mln"]))
    )
    cmp["flag_fx_change"] = existing & cmp["fx_rel_change"].abs().ge(float(thresholds["fx_relative_change"]))
    cmp["flag_reserve_jump"] = cmp["delta_reserve_mln"].abs().ge(float(thresholds["reserve_jump_mln"]))

    # Heuristic primary driver. Every contract has exactly one category; sums are additive.
    conditions = [
        cmp["movement_type"].eq("NEW"),
        cmp["movement_type"].eq("EXIT"),
        cmp["flag_new_default"],
        cmp["flag_stage_deterioration"],
        cmp["flag_stage_improvement"],
        cmp["flag_rating_downgrade"],
        cmp["flag_overdue_deterioration"],
        cmp["flag_overdue_improvement"],
        cmp["flag_overdue_change"],
        cmp["flag_pd_change"],
        cmp["flag_lgd_change"],
        cmp["flag_ead_change"],
        cmp["flag_fx_change"],
    ]
    choices = [
        "NEW", "EXIT", "NEW_DEFAULT", "STAGE_DETERIORATION", "STAGE_IMPROVEMENT", "RATING_DOWNGRADE",
        "OVERDUE_DETERIORATION", "OVERDUE_IMPROVEMENT", "OVERDUE_CHANGE", "PD_CHANGE", "LGD_CHANGE", "EAD_CHANGE", "FX_CHANGE",
    ]
    cmp["primary_driver"] = np.select(conditions, choices, default="OTHER")

    flag_cols = [
        "flag_new_default", "flag_stage_change", "flag_rating_change", "flag_overdue_change",
        "flag_overdue_deterioration", "flag_overdue_improvement", "flag_pd_change",
        "flag_lgd_change", "flag_ead_change", "flag_fx_change",
    ]
    cmp["material_factor_count"] = cmp[flag_cols].sum(axis=1)
    cmp["multi_factor"] = existing & cmp["material_factor_count"].gt(1)
    cmp["driver_flags"] = cmp.apply(
        lambda row: ";".join([col.replace("flag_", "").upper() for col in flag_cols if bool(row[col])]) or "NONE",
        axis=1,
    )

    # Optional exact FX effect when original money is in transaction currency.
    if not bool(cfg["input"].get("monetary_values_in_byn", True)):
        q_prev = pd.to_numeric(cmp.get("reserve_original_prev"), errors="coerce").fillna(0)
        q_curr = pd.to_numeric(cmp.get("reserve_original_curr"), errors="coerce").fillna(0)
        f_prev = fx_prev.fillna(1)
        f_curr = fx_curr.fillna(f_prev).fillna(1)
        cmp["fx_effect_mln"] = np.where(existing, q_prev * (f_curr - f_prev) / 1_000_000, 0.0)
        cmp["local_amount_effect_mln"] = np.where(existing, (q_curr - q_prev) * f_curr / 1_000_000, 0.0)
    else:
        cmp["fx_effect_mln"] = np.nan
        cmp["local_amount_effect_mln"] = np.nan

    # Coalesced identifying/dimensional fields for Qlik/top movements.
    carry = [
        "crm_id", "client_name", "source", "contract_number", "tranche_number", "industry",
        "subportfolio", "business_segment", "risk_segment", "currency_norm",
    ]
    for base in carry:
        curr_col, prev_col = f"{base}_curr", f"{base}_prev"
        if curr_col in cmp and prev_col in cmp:
            cmp[base] = cmp[curr_col].combine_first(cmp[prev_col])
        elif curr_col in cmp:
            cmp[base] = cmp[curr_col]
        elif prev_col in cmp:
            cmp[base] = cmp[prev_col]

    preferred = [
        "report_date", "previous_report_date", "contract_key", "crm_id", "client_name", "source",
        "contract_number", "tranche_number", "industry", "subportfolio", "business_segment",
        "risk_segment", "currency_norm", "movement_type", "primary_driver", "driver_flags",
        "multi_factor", "reserve_mln_prev", "reserve_mln_curr", "delta_reserve_mln",
        "ead_mln_prev", "ead_mln_curr", "delta_ead_mln", "stage_prev", "stage_curr",
        "stage_migration", "rating_prev", "rating_curr", "rating_migration", "overdue_prev",
        "overdue_curr", "pd_pit_prev", "pd_pit_curr", "pd_rel_change", "lgd_prev", "lgd_curr",
        "lgd_abs_change", "fx_rate_prev", "fx_rate_curr", "fx_effect_mln",
        "local_amount_effect_mln", "flag_stage_change", "flag_stage_deterioration",
        "flag_stage_improvement", "flag_rating_change", "flag_rating_downgrade",
        "flag_new_default", "flag_overdue_change", "flag_overdue_deterioration", "flag_overdue_improvement", "flag_pd_change", "flag_lgd_change", "flag_ead_change",
        "flag_fx_change", "flag_reserve_jump",
    ]
    return cmp[[x for x in preferred if x in cmp.columns]].copy()


def create_dq_report(df: pd.DataFrame, report_date: pd.Timestamp, cfg: dict[str, Any]) -> tuple[dict[str, Any], pd.DataFrame]:
    c = cfg["columns"]
    reserve_gt_ead = (df["reserve_byn"].abs() > df["ead_byn"].abs()) & df["ead_byn"].notna()
    pd_series = pd.to_numeric(df.get("pd_pit"), errors="coerce")
    lgd_series = pd.to_numeric(df.get("lgd"), errors="coerce")
    metrics = {
        "report_date": report_date.strftime("%Y-%m-%d"),
        "rows": int(len(df)),
        "contracts": int(df["contract_key"].nunique()),
        "clients": int(df.get("crm_id", pd.Series(dtype=str)).nunique()),
        "duplicate_keys": int(df["contract_key"].duplicated().sum()),
        "missing_crm_id": int(df.get("crm_id", pd.Series(dtype=str)).isna().sum()),
        "missing_contract_number": int(df.get("contract_number", pd.Series(dtype=str)).isna().sum()),
        "unknown_stage": int(df["stage_norm"].eq("UNKNOWN").sum()),
        "missing_ead": int(df[c["ead"]].isna().sum()),
        "negative_ead": int(df["ead_byn"].lt(0).sum()),
        "negative_reserve": int(df["reserve_byn"].lt(0).sum()),
        "reserve_gt_ead": int(reserve_gt_ead.sum()),
        "pd_out_of_range": int(((pd_series < 0) | (pd_series > 1)).fillna(False).sum()),
        "lgd_out_of_range": int(((lgd_series < 0) | (lgd_series > 1)).fillna(False).sum()),
    }
    hard_errors = metrics["duplicate_keys"] + metrics["missing_contract_number"]
    metrics["status"] = "FAIL" if hard_errors > 0 else ("WARN" if sum([
        metrics["unknown_stage"], metrics["missing_ead"], metrics["pd_out_of_range"], metrics["lgd_out_of_range"]
    ]) > 0 else "OK")
    detail = pd.DataFrame([{"metric": k, "value": v} for k, v in metrics.items()])
    return metrics, detail


def daily_summary(df: pd.DataFrame, movements: pd.DataFrame | None, report_date: pd.Timestamp) -> pd.DataFrame:
    reserve_total = float(df["reserve_mln"].sum())
    ead_total = float(df["ead_mln"].sum())
    row: dict[str, Any] = {
        "report_date": report_date.strftime("%Y-%m-%d"),
        "ead_mln": ead_total,
        "reserve_stage1_mln": float(df.loc[df["stage_norm"].eq("1"), "reserve_mln"].sum()),
        "reserve_stage2_mln": float(df.loc[df["stage_norm"].eq("2"), "reserve_mln"].sum()),
        "reserve_stage3_mln": float(df.loc[df["stage_norm"].eq("3"), "reserve_mln"].sum()),
        "reserve_unknown_stage_mln": float(df.loc[df["stage_norm"].eq("UNKNOWN"), "reserve_mln"].sum()),
        "reserve_individual_mln": float(df.loc[df["individual_flag"].fillna(False), "reserve_mln"].sum()),
        "reserve_overdue_mln": float(df["reserve_overdue_mln"].sum()),
        "reserve_total_mln": reserve_total,
        "reserve_rate": reserve_total / ead_total if ead_total else np.nan,
        "contracts": int(df["contract_key"].nunique()),
        "clients": int(df.get("crm_id", pd.Series(dtype=str)).nunique()),
    }
    if movements is None:
        row.update({
            "new_effect_mln": 0.0, "exit_effect_mln": 0.0, "existing_effect_mln": 0.0,
            "delta_reserve_dod_mln": np.nan, "delta_ead_dod_mln": np.nan, "bridge_check_mln": np.nan,
        })
    else:
        new_effect = float(movements.loc[movements["movement_type"].eq("NEW"), "delta_reserve_mln"].sum())
        exit_effect = float(movements.loc[movements["movement_type"].eq("EXIT"), "delta_reserve_mln"].sum())
        existing_effect = float(movements.loc[movements["movement_type"].eq("EXISTING"), "delta_reserve_mln"].sum())
        delta = float(movements["delta_reserve_mln"].sum())
        delta_ead = float(movements["delta_ead_mln"].sum())
        row.update({
            "new_effect_mln": new_effect,
            "exit_effect_mln": exit_effect,
            "existing_effect_mln": existing_effect,
            "delta_reserve_dod_mln": delta,
            "delta_ead_dod_mln": delta_ead,
            "bridge_check_mln": delta - new_effect - exit_effect - existing_effect,
        })
    return pd.DataFrame([row])


def build_drivers(df: pd.DataFrame, movements: pd.DataFrame | None, cfg: dict[str, Any], report_date: pd.Timestamp) -> pd.DataFrame:
    c = cfg["columns"]
    frames: list[pd.DataFrame] = []
    dim_alias = {
        c["stage"]: "stage_norm",
        c["rating"]: "rating_norm",
        c["overdue_bucket"]: "overdue_bucket_norm",
        c["currency"]: "currency_norm",
        c["risk_segment"]: "risk_segment",
        c["business_segment"]: "business_segment",
        c["subportfolio"]: "subportfolio",
        c["industry"]: "industry",
    }
    for configured in cfg["monitor_dimensions"]:
        dim = dim_alias.get(configured, configured)
        if dim not in df.columns:
            continue
        tmp = df.copy()
        tmp[dim] = safe_str(tmp[dim], missing="UNKNOWN")
        agg = tmp.groupby(dim, dropna=False).agg(
            ead_mln=("ead_mln", "sum"), reserve_mln=("reserve_mln", "sum"),
            contracts=("contract_key", "nunique"), clients=("crm_id", "nunique"),
        ).reset_index().rename(columns={dim: "dimension_value"})
        agg["dimension"] = configured
        agg["report_date"] = report_date.strftime("%Y-%m-%d")
        agg["reserve_rate"] = np.where(agg["ead_mln"].ne(0), agg["reserve_mln"] / agg["ead_mln"], np.nan)
        frames.append(agg)

    result = pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()

    # Add exact primary-driver effect rows from movements.
    if movements is not None and not movements.empty:
        drv = movements.groupby("primary_driver", dropna=False).agg(
            delta_reserve_mln=("delta_reserve_mln", "sum"),
            delta_ead_mln=("delta_ead_mln", "sum"),
            contracts=("contract_key", "nunique"),
        ).reset_index().rename(columns={"primary_driver": "dimension_value"})
        drv["dimension"] = "PRIMARY_DRIVER"
        drv["report_date"] = report_date.strftime("%Y-%m-%d")
        drv["ead_mln"] = np.nan
        drv["reserve_mln"] = np.nan
        drv["reserve_rate"] = np.nan
        drv["clients"] = np.nan
        result = pd.concat([result, drv], ignore_index=True, sort=False)
    return result


def build_top_movements(movements: pd.DataFrame | None, n: int) -> pd.DataFrame:
    if movements is None or movements.empty:
        return pd.DataFrame()
    out = movements.copy()
    out["abs_delta_reserve_mln"] = out["delta_reserve_mln"].abs()
    out = out.sort_values("abs_delta_reserve_mln", ascending=False).head(n)
    return out


def upsert_csv(path: Path, new_df: pd.DataFrame, key_columns: list[str], replace_filter: dict[str, Any] | None = None) -> None:
    if new_df is None or new_df.empty:
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        old = pd.read_csv(path)
        if replace_filter:
            mask = pd.Series(True, index=old.index)
            for col, value in replace_filter.items():
                if col in old.columns:
                    mask &= old[col].astype(str).eq(str(value))
            old = old.loc[~mask]
        combined = pd.concat([old, new_df], ignore_index=True, sort=False)
    else:
        combined = new_df.copy()
    existing_keys = [k for k in key_columns if k in combined.columns]
    if existing_keys:
        combined = combined.drop_duplicates(existing_keys, keep="last")
    combined.to_csv(path, index=False, encoding="utf-8-sig")


def apply_mtd(summary_path: Path) -> None:
    if not summary_path.exists():
        return
    df = pd.read_csv(summary_path)
    if df.empty:
        return
    df["report_date"] = pd.to_datetime(df["report_date"])
    df = df.sort_values("report_date")
    df["month"] = df["report_date"].dt.to_period("M")
    df["month_start_reserve_mln"] = df.groupby("month")["reserve_total_mln"].transform("first")
    df["delta_reserve_mtd_mln"] = df["reserve_total_mln"] - df["month_start_reserve_mln"]
    df["month_start_ead_mln"] = df.groupby("month")["ead_mln"].transform("first")
    df["delta_ead_mtd_mln"] = df["ead_mln"] - df["month_start_ead_mln"]
    df = df.drop(columns=["month"])
    df["report_date"] = df["report_date"].dt.strftime("%Y-%m-%d")
    df.to_csv(summary_path, index=False, encoding="utf-8-sig")


def write_dq_files(ctx: RunContext, metrics: dict[str, Any], detail: pd.DataFrame, report_date: pd.Timestamp) -> None:
    stem = report_date.strftime("%Y-%m-%d")
    with (ctx.qlik_dir / f"dq_{stem}.json").open("w", encoding="utf-8") as f:
        json.dump(metrics, f, ensure_ascii=False, indent=2)
    detail.to_csv(ctx.qlik_dir / f"dq_{stem}.csv", index=False, encoding="utf-8-sig")
    upsert_csv(
        ctx.qlik_dir / "dq_history.csv",
        pd.DataFrame([metrics]),
        key_columns=["report_date"],
        replace_filter={"report_date": stem},
    )


def build_daily_excel(
    ctx: RunContext,
    report_date: pd.Timestamp,
    summary: pd.DataFrame,
    movements: pd.DataFrame | None,
    drivers: pd.DataFrame,
    top: pd.DataFrame,
    dq: dict[str, Any],
) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Daily_Monitoring"
    dark = "1F4E78"
    blue = "4472C4"
    green = "70AD47"
    yellow = "FFF2CC"
    red = "F4CCCC"
    white = "FFFFFF"

    ws["A1"] = f"Мониторинг резервов на {report_date.strftime('%d.%m.%Y')}"
    ws["A1"].font = Font(size=16, bold=True, color=white)
    ws["A1"].fill = PatternFill("solid", fgColor=dark)
    ws.merge_cells("A1:H1")

    s = summary.iloc[0]
    kpis = [
        ("Резерв, млн BYN", s.get("reserve_total_mln")),
        ("EAD, млн BYN", s.get("ead_mln")),
        ("Reserve rate", s.get("reserve_rate")),
        ("Δ резерв DoD, млн", s.get("delta_reserve_dod_mln")),
        ("Δ резерв MTD, млн", s.get("delta_reserve_mtd_mln", np.nan)),
        ("Stage 2+3, млн", s.get("reserve_stage2_mln", 0) + s.get("reserve_stage3_mln", 0)),
        ("Договоров", s.get("contracts")),
        ("DQ status", dq.get("status")),
    ]
    for col, (label, value) in enumerate(kpis, start=1):
        cell = ws.cell(row=3, column=col, value=label)
        cell.font = Font(bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.alignment = Alignment(wrap_text=True, horizontal="center")
        ws.cell(row=4, column=col, value=None if pd.isna(value) else value)
    ws["C4"].number_format = "0.00%"
    for cell in ws[4]:
        if cell.column != 3 and isinstance(cell.value, (float, np.floating)):
            cell.number_format = "#,##0.00"

    # Exact top-level bridge.
    ws["A7"] = "Мост изменения резерва"
    ws["A7"].font = Font(bold=True, color=white)
    ws["A7"].fill = PatternFill("solid", fgColor=green)
    bridge = [
        ("NEW", s.get("new_effect_mln")),
        ("EXIT", s.get("exit_effect_mln")),
        ("EXISTING", s.get("existing_effect_mln")),
        ("ИТОГО Δ", s.get("delta_reserve_dod_mln")),
        ("Проверка", s.get("bridge_check_mln")),
    ]
    for i, (name, value) in enumerate(bridge, start=8):
        ws.cell(i, 1, name)
        ws.cell(i, 2, None if pd.isna(value) else value).number_format = "#,##0.00"

    # Primary-driver bridge.
    ws["D7"] = "Primary driver внутри движения"
    ws["D7"].font = Font(bold=True, color=white)
    ws["D7"].fill = PatternFill("solid", fgColor=green)
    if movements is not None and not movements.empty:
        drv = movements.groupby("primary_driver", as_index=False)["delta_reserve_mln"].sum().sort_values("delta_reserve_mln", key=abs, ascending=False)
        for i, row in enumerate(drv.itertuples(index=False), start=8):
            ws.cell(i, 4, row.primary_driver)
            ws.cell(i, 5, row.delta_reserve_mln).number_format = "#,##0.00"

    # Top movements sheet.
    wt = wb.create_sheet("Top_Movements")
    top_cols = [
        "client_name", "contract_number", "tranche_number", "movement_type", "primary_driver",
        "reserve_mln_prev", "reserve_mln_curr", "delta_reserve_mln", "ead_mln_prev", "ead_mln_curr",
        "stage_prev", "stage_curr", "rating_prev", "rating_curr", "pd_pit_prev", "pd_pit_curr",
        "lgd_prev", "lgd_curr", "driver_flags",
    ]
    top_view = top[[c for c in top_cols if c in top.columns]].copy() if top is not None and not top.empty else pd.DataFrame(columns=top_cols)
    for j, col in enumerate(top_view.columns, 1):
        wt.cell(1, j, col)
        wt.cell(1, j).font = Font(bold=True, color=white)
        wt.cell(1, j).fill = PatternFill("solid", fgColor=blue)
    for i, row in enumerate(top_view.itertuples(index=False), 2):
        for j, value in enumerate(row, 1):
            wt.cell(i, j, None if pd.isna(value) else value)
    if "delta_reserve_mln" in top_view.columns and len(top_view) > 0:
        idx = list(top_view.columns).index("delta_reserve_mln") + 1
        rng = f"{get_column_letter(idx)}2:{get_column_letter(idx)}{len(top_view)+1}"
        wt.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color="F4CCCC", mid_type="percentile", mid_value=50, mid_color="FFFFFF", end_type="max", end_color="D9EAD3"))

    # Drivers sheet.
    wd = wb.create_sheet("Drivers")
    for j, col in enumerate(drivers.columns, 1):
        wd.cell(1, j, col)
        wd.cell(1, j).font = Font(bold=True, color=white)
        wd.cell(1, j).fill = PatternFill("solid", fgColor=blue)
    for i, row in enumerate(drivers.itertuples(index=False), 2):
        for j, value in enumerate(row, 1):
            wd.cell(i, j, None if pd.isna(value) else value)

    # Data quality sheet.
    wq = wb.create_sheet("Data_Quality")
    wq.append(["metric", "value"])
    for cell in wq[1]:
        cell.font = Font(bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=blue)
    for k, v in dq.items():
        wq.append([k, v])

    # Trend sheet from cumulative summary.
    summary_path = ctx.qlik_dir / "daily_summary.csv"
    if summary_path.exists():
        hist = pd.read_csv(summary_path)
        wh = wb.create_sheet("History")
        for j, col in enumerate(hist.columns, 1):
            wh.cell(1, j, col)
            wh.cell(1, j).font = Font(bold=True, color=white)
            wh.cell(1, j).fill = PatternFill("solid", fgColor=blue)
        for i, row in enumerate(hist.itertuples(index=False), 2):
            for j, value in enumerate(row, 1):
                wh.cell(i, j, None if pd.isna(value) else value)
        if len(hist) >= 2 and "reserve_total_mln" in hist.columns:
            chart = LineChart()
            chart.title = "Динамика резерва"
            date_col = list(hist.columns).index("report_date") + 1
            value_col = list(hist.columns).index("reserve_total_mln") + 1
            data = Reference(wh, min_col=value_col, min_row=1, max_row=len(hist)+1)
            cats = Reference(wh, min_col=date_col, min_row=2, max_row=len(hist)+1)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            chart.height = 7
            chart.width = 14
            ws.add_chart(chart, "A16")

    # Formatting.
    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2" if sheet.title != "Daily_Monitoring" else "A3"
        for col in range(1, sheet.max_column + 1):
            max_len = 0
            for row in range(1, min(sheet.max_row, 100) + 1):
                value = sheet.cell(row, col).value
                if value is not None:
                    max_len = max(max_len, len(str(value)))
            sheet.column_dimensions[get_column_letter(col)].width = min(max(max_len + 2, 11), 35)
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    out = ctx.reports_dir / f"{report_date.strftime('%Y-%m-%d')}_reserve_monitoring.xlsx"
    wb.save(out)
    return out


def reconcile(summary: pd.DataFrame, movements: pd.DataFrame | None, prev: pd.DataFrame | None) -> None:
    if movements is None or prev is None:
        return
    actual = float(summary.iloc[0]["delta_reserve_dod_mln"])
    expected = float(summary.iloc[0]["reserve_total_mln"] - prev["reserve_mln"].sum())
    if not np.isclose(actual, expected, atol=1e-7, rtol=1e-9):
        raise RuntimeError(f"Reserve reconciliation failed: movement sum={actual}, snapshot delta={expected}")
    bridge = float(summary.iloc[0]["bridge_check_mln"])
    if not np.isclose(bridge, 0.0, atol=1e-7):
        raise RuntimeError(f"Bridge reconciliation failed: {bridge}")


def process_file(ctx: RunContext, input_path: Path) -> dict[str, Path]:
    cfg = ctx.cfg
    logging.info("Reading snapshot: %s", input_path)
    raw = read_input(input_path, sheet_name=cfg["input"].get("sheet_name", 0))
    snapshot, report_date = prepare_snapshot(raw, cfg)
    contract_daily = select_contract_daily_columns(snapshot)

    date_str = report_date.strftime("%Y-%m-%d")
    daily_csv = ctx.processed_daily_dir / f"{date_str}.csv"
    daily_parquet = ctx.processed_daily_dir / f"{date_str}.parquet"
    # Preserve the full normalized source snapshot plus derived fields for audit/reprocessing.
    save_frame(snapshot, daily_csv, daily_parquet)

    dq_metrics, dq_detail = create_dq_report(snapshot, report_date, cfg)
    write_dq_files(ctx, dq_metrics, dq_detail, report_date)
    logging.info("Data quality status: %s", dq_metrics["status"])

    prev_path = previous_daily_file(ctx, report_date)
    if prev_path:
        prev = pd.read_parquet(prev_path) if prev_path.suffix.lower() == ".parquet" else pd.read_csv(prev_path)
        if "report_date" in prev.columns:
            prev["report_date"] = pd.to_datetime(prev["report_date"], errors="coerce")
    else:
        prev = None
    movements = None
    movement_csv = None
    movement_parquet = None
    if prev is not None:
        prev_date = pd.to_datetime(prev["report_date"]).max().normalize()
        logging.info("Comparing with previous snapshot: %s", prev_date.date())
        movements = build_movements(snapshot, prev, cfg)
        movement_stem = f"{date_str}_vs_{prev_date.strftime('%Y-%m-%d')}"
        movement_csv = ctx.processed_movements_dir / f"{movement_stem}.csv"
        movement_parquet = ctx.processed_movements_dir / f"{movement_stem}.parquet"
        save_frame(movements, movement_csv, movement_parquet)
    else:
        logging.info("No previous snapshot found. This is the baseline day.")

    summary = daily_summary(snapshot, movements, report_date)
    summary_path = ctx.qlik_dir / "daily_summary.csv"
    upsert_csv(summary_path, summary, ["report_date"], {"report_date": date_str})
    apply_mtd(summary_path)
    summary_all = pd.read_csv(summary_path)
    summary_today = summary_all.loc[summary_all["report_date"].astype(str).eq(date_str)].copy()

    drivers = build_drivers(snapshot, movements, cfg, report_date)
    upsert_csv(
        ctx.qlik_dir / "daily_drivers.csv",
        drivers,
        ["report_date", "dimension", "dimension_value"],
        {"report_date": date_str},
    )

    top = build_top_movements(movements, int(cfg["project"].get("top_n_movements", 50)))
    if movements is not None:
        upsert_csv(
            ctx.qlik_dir / "movements.csv",
            movements,
            ["report_date", "contract_key"],
            {"report_date": date_str},
        )
    if not top.empty:
        upsert_csv(
            ctx.qlik_dir / "top_movements.csv",
            top,
            ["report_date", "contract_key"],
            {"report_date": date_str},
        )

    # Latest contract state for simple Qlik client drill-down; full history is loaded from daily folder.
    contract_daily.to_csv(ctx.qlik_dir / "latest_contracts.csv", index=False, encoding="utf-8-sig")
    reconcile(summary_today, movements, prev)

    outputs: dict[str, Path] = {
        "daily_csv": daily_csv,
        "summary_csv": summary_path,
        "drivers_csv": ctx.qlik_dir / "daily_drivers.csv",
        "dq_history": ctx.qlik_dir / "dq_history.csv",
    }
    if daily_parquet.exists():
        outputs["daily_parquet"] = daily_parquet
    if movement_csv:
        outputs["movement_csv"] = movement_csv
    if movement_parquet and movement_parquet.exists():
        outputs["movement_parquet"] = movement_parquet

    if bool(cfg["project"].get("export_daily_excel", True)):
        report_path = build_daily_excel(ctx, report_date, summary_today, movements, drivers, top, dq_metrics)
        outputs["daily_excel"] = report_path

    logging.info("Finished report date %s", date_str)
    return outputs


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Daily reserve monitoring")
    parser.add_argument("--config", default="config.yaml", help="Path to config.yaml")
    parser.add_argument("--file", default=None, help="Raw snapshot file. If omitted, latest file from raw_dir is used.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    ctx = resolve_context(Path(args.config))
    setup_logging(ctx)
    try:
        input_path = Path(args.file).resolve() if args.file else find_latest_raw(ctx)
        outputs = process_file(ctx, input_path)
        print("\nCreated/updated:")
        for name, path in outputs.items():
            print(f"- {name}: {path}")
        return 0
    except Exception:
        logging.exception("Monitoring failed")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
